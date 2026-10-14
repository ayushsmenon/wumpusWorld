import shutil
import subprocess
import tempfile
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ingestion.models import (CorruptFileError, EmptyFileError, ParsedUnit,
                              PasswordProtectedError)
from ingestion.pdf_parser import render_page

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
MIN_TEXT_CHARS = 25
BIG_PICTURE_FRACTION = 0.05
DRAWN_SHAPES_FOR_DIAGRAM = 5
SCAN_PICTURE_FRACTION = 0.5
DRAWN_TYPES = {MSO_SHAPE_TYPE.AUTO_SHAPE,
               MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.LINE}
LIBREOFFICE_TIMEOUT = 300


def safe_type(shape):
    try:
        return shape.shape_type
    except Exception:
        return None


def walk(shapes):
    for shape in shapes:
        yield shape
        if safe_type(shape) == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes)


def save_picture(shape, slide_number, index, out_dir: Path) -> dict:
    path, size = None, (None, None)
    try:
        image = shape.image
        out_dir.mkdir(parents=True, exist_ok=True)
        filename = f"slide_{slide_number:03d}_img_{index}.{image.ext}"
        (out_dir / filename).write_bytes(image.blob)
        path = f"{out_dir.name}/{filename}"
        size = image.size
    except Exception:
        pass
    return {"kind": "picture", "path": path, "bbox": None, "width": size[0], "height": size[1]}


def slide_title(slide) -> str:
    try:
        if slide.shapes.title is not None:
            return slide.shapes.title.text_frame.text.strip()
    except Exception:
        pass
    return ""


def parse_one_slide(slide, number, slide_area, out_dir: Path) -> ParsedUnit:
    text_parts, tables, images = [], [], []
    drawn = 0
    biggest_picture = 0

    reading_order = sorted(
        slide.shapes, key=lambda s: (s.top or 0, s.left or 0))
    for shape in walk(reading_order):
        kind = safe_type(shape)
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                text_parts.append(t)
        if shape.has_table:
            rows = [[cell.text.replace("\n", " ").strip() for cell in row.cells]
                    for row in shape.table.rows]
            if rows:
                tables.append(rows)
                text_parts.append("\n".join(" | ".join(r) for r in rows))
        if shape.has_chart:
            images.append({"kind": "chart", "path": None, "bbox": None,
                           "width": None, "height": None})
        if kind == MSO_SHAPE_TYPE.PICTURE:
            area = (shape.width or 0) * (shape.height or 0)
            if area > BIG_PICTURE_FRACTION * slide_area:
                biggest_picture = max(biggest_picture, area)
                images.append(save_picture(
                    shape, number, len(images) + 1, out_dir))
        elif kind in DRAWN_TYPES:
            drawn += 1

    text = "\n".join(text_parts)
    notes = ""
    if slide.has_notes_slide:
        frame = slide.notes_slide.notes_text_frame
        notes = frame.text.strip() if frame is not None else ""

    has_visuals = bool(images) or drawn >= DRAWN_SHAPES_FOR_DIAGRAM
    needs_ocr = len(text) < MIN_TEXT_CHARS \
        and biggest_picture > SCAN_PICTURE_FRACTION * slide_area
    return ParsedUnit("slide", number, text, tables, images, has_visuals, needs_ocr,
                      notes=notes,
                      metadata={"title": slide_title(slide), "drawn_shapes": drawn})


def find_soffice():
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for candidate in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                      r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if Path(candidate).exists():
            return candidate
    return None


def render_slides(path: Path, slide_numbers, total_slides, out_dir: Path) -> dict:

    soffice = find_soffice()
    if not soffice or not slide_numbers:
        return {}
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(path)],
                           check=True, timeout=LIBREOFFICE_TIMEOUT, capture_output=True)
            pdfs = list(Path(tmp).glob("*.pdf"))
            if not pdfs:
                return {}
            doc = pymupdf.open(pdfs[0])
        except Exception:
            return {}
        try:
            if doc.page_count != total_slides:   # hidden slides etc. would misalign numbers
                return {}
            return {n: render_page(doc[n - 1], n, out_dir, prefix="slide_render")
                    for n in slide_numbers}
        finally:
            doc.close()


def parse_slides(path: Path, out_dir: Path) -> list:
    with open(path, "rb") as f:
        header = f.read(8)
    if header == OLE_MAGIC:
        raise PasswordProtectedError(
            "Presentation is password protected (or an old .ppt)")

    try:
        prs = Presentation(str(path))
    except Exception as e:
        raise CorruptFileError(f"Could not open slides: {e}") from e

    if len(prs.slides) == 0:
        raise EmptyFileError("Presentation has no slides")

    slide_area = (prs.slide_width or 1) * (prs.slide_height or 1)
    units = [parse_one_slide(slide, n, slide_area, out_dir)
             for n, slide in enumerate(prs.slides, start=1)]

    flagged = [u.number for u in units if u.has_visuals]
    renders = render_slides(path, flagged, len(units), out_dir)
    for u in units:
        if u.number in renders:
            u.images.append(renders[u.number])
        elif u.has_visuals:
            u.metadata["render_missing"] = True
    return units
